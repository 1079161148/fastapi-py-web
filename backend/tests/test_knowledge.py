"""知识库入库 / 检索 / RAG 集成测试（真实数据库 + 假嵌入、假重排）。

覆盖设计文档 M3 验收点：
- 上传 → 入队 → worker 解析分块向量化 → 状态 done、chunk_count 正确
- 混合检索（向量 + 关键词 OR + RRF）能命中相关段落
- rerank 生效、降级不影响回答
- 对话 SSE 含 sources 事件，assistant 消息固化引用来源
- 租户隔离（RLS）与角色权限（非管理员不可写知识库）
- 分块与解析的纯函数行为（标题元数据、窗口重叠、DOCX 抽取）
"""

import io
import uuid

import pytest
from chat_helpers import PASSWORD, auth_headers, parse_sse, register_user
from httpx import AsyncClient
from knowledge_helpers import (
    FakeEmbedder,
    FakeQueue,
    FakeReranker,
    create_kb,
    upload_markdown,
)

from app.core.config import get_settings
from app.core.security import hash_password
from app.core.tenant import set_current_tenant
from app.db.models.knowledge import CHUNK_STRATEGY_BY_HEADING, CHUNK_STRATEGY_FIXED_WINDOW
from app.db.models.tenant import TenantMember
from app.db.models.user import User
from app.db.session import async_session_factory
from app.services.ingestion import (
    Section,
    chunk_sections,
    extract_sections,
    tokenize_for_search,
)
from app.workers.tasks import process_document


def _patch_worker_embedder(monkeypatch: pytest.MonkeyPatch, embedder: FakeEmbedder) -> None:
    """worker 直接 import 了工厂函数，必须替换模块属性而不是依赖覆盖。"""
    monkeypatch.setattr("app.workers.tasks.get_embedding_client", lambda: embedder)


async def test_upload_process_and_search(
    client: AsyncClient,
    unique_email: str,
    fake_queue: FakeQueue,
    fake_embedder: FakeEmbedder,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = await register_user(client, unique_email)
    tenant_id = data["user"]["tenant_id"]
    headers = auth_headers(data["tokens"]["access_token"])

    kb = await create_kb(client, headers)
    document = await upload_markdown(client, headers, kb["id"])

    assert document["status"] == "uploaded"
    assert fake_queue.jobs and fake_queue.jobs[0][0] == "process_document"

    _patch_worker_embedder(monkeypatch, fake_embedder)
    await process_document({}, document["id"], tenant_id)

    detail = await client.get(f"/api/v1/knowledge/documents/{document['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["status"] == "done", detail.json()["error_msg"]
    assert detail.json()["chunk_count"] >= 2

    bases = await client.get("/api/v1/knowledge/bases", headers=headers)
    assert bases.json()[0]["document_count"] == 1

    hits = await client.get(
        "/api/v1/knowledge/search", params={"q": "退款多久到账"}, headers=headers
    )
    assert hits.status_code == 200, hits.text
    items = hits.json()["items"]
    assert items, "向量/关键词至少有一路应命中退款段落"
    assert any("退款" in item["content"] for item in items)
    assert any(item["keyword_score"] is not None for item in items), "关键词通道应当有命中"


async def test_chat_emits_sources_and_persists_them(
    client: AsyncClient,
    unique_email: str,
    fake_queue: FakeQueue,
    fake_embedder: FakeEmbedder,
    fake_llm,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = await register_user(client, unique_email)
    tenant_id = data["user"]["tenant_id"]
    headers = auth_headers(data["tokens"]["access_token"])

    kb = await create_kb(client, headers)
    document = await upload_markdown(client, headers, kb["id"])
    _patch_worker_embedder(monkeypatch, fake_embedder)
    await process_document({}, document["id"], tenant_id)

    conversation = await client.post("/api/v1/conversations", json={}, headers=headers)
    conversation_id = conversation.json()["id"]

    resp = await client.post(
        f"/api/v1/conversations/{conversation_id}/messages/stream",
        json={"content": "退款多久到账？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    events = parse_sse(resp.text)
    names = [name for name, _ in events]
    assert names[0] == "message"
    assert names[1] == "sources", f"引用应先行推送，实际顺序：{names}"

    sources_event = events[1][1]
    assert sources_event["degraded"] is False
    assert sources_event["items"]
    assert sources_event["items"][0]["document_name"] == "policy.md"

    # 参考资料进入 system prompt
    system_content = fake_llm.last_prompt[0].content
    assert "参考资料" in system_content
    assert "[1]" in system_content

    messages = await client.get(
        f"/api/v1/conversations/{conversation_id}/messages", headers=headers
    )
    assistant = [item for item in messages.json()["items"] if item["role"] == "assistant"]
    assert len(assistant) == 1
    assert assistant[0]["sources"], "assistant 消息必须固化引用来源"
    assert assistant[0]["sources"][0]["chunk_id"] == sources_event["items"][0]["chunk_id"]


async def test_chat_degrades_when_embedding_missing(
    client: AsyncClient,
    unique_email: str,
    fake_llm,
    fake_queue: FakeQueue,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """嵌入模型未配置：有知识库也照常回答，但必须显式标注"未引用"。"""
    settings = get_settings()
    monkeypatch.setattr(settings, "ai_embedding_model", "")
    monkeypatch.setattr(settings, "ai_embedding_api_key", "")
    monkeypatch.setattr(settings, "ai_embedding_base_url", "")
    monkeypatch.setattr(settings, "ai_api_key", "")
    monkeypatch.setattr(settings, "ai_base_url", "")

    data = await register_user(client, unique_email)
    headers = auth_headers(data["tokens"]["access_token"])
    await create_kb(client, headers)

    conversation_id = (await client.post("/api/v1/conversations", json={}, headers=headers)).json()[
        "id"
    ]
    resp = await client.post(
        f"/api/v1/conversations/{conversation_id}/messages/stream",
        json={"content": "退款多久到账？"},
        headers=headers,
    )
    assert resp.status_code == 200

    events = parse_sse(resp.text)
    sources_event = next(data for name, data in events if name == "sources")
    assert sources_event["degraded"] is True
    assert "嵌入模型" in sources_event["reason"]
    assert sources_event["items"] == []
    # 回答本身仍然完成
    assert events[-1][0] == "done"
    # 无依据时 prompt 不应包含参考资料段
    assert "参考资料" not in fake_llm.last_prompt[0].content


async def test_rerank_is_applied(
    client: AsyncClient,
    unique_email: str,
    fake_queue: FakeQueue,
    fake_embedder: FakeEmbedder,
    fake_reranker: FakeReranker,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = await register_user(client, unique_email)
    tenant_id = data["user"]["tenant_id"]
    headers = auth_headers(data["tokens"]["access_token"])
    kb = await create_kb(client, headers)
    document = await upload_markdown(client, headers, kb["id"])
    _patch_worker_embedder(monkeypatch, fake_embedder)
    await process_document({}, document["id"], tenant_id)

    hits = await client.get("/api/v1/knowledge/search", params={"q": "退款流程"}, headers=headers)
    items = hits.json()["items"]
    assert fake_reranker.calls == 1
    assert items and all(item["reranked"] for item in items)


async def test_upload_rejects_bad_input(
    client: AsyncClient, unique_email: str, fake_queue: FakeQueue, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = await register_user(client, unique_email)
    headers = auth_headers(data["tokens"]["access_token"])
    kb = await create_kb(client, headers)

    bad_type = await client.post(
        "/api/v1/knowledge/documents",
        data={"kb_id": kb["id"]},
        files={"file": ("virus.exe", b"MZ...", "application/octet-stream")},
        headers=headers,
    )
    assert bad_type.status_code == 400
    assert bad_type.json()["code"] == "UNSUPPORTED_FILE_TYPE"

    monkeypatch.setattr(get_settings(), "upload_max_mb", 0)
    too_large = await client.post(
        "/api/v1/knowledge/documents",
        data={"kb_id": kb["id"]},
        files={"file": ("big.md", b"# x" * 100, "text/markdown")},
        headers=headers,
    )
    assert too_large.status_code == 413
    assert too_large.json()["code"] == "FILE_TOO_LARGE"


async def test_knowledge_is_tenant_isolated(
    client: AsyncClient,
    unique_email: str,
    fake_queue: FakeQueue,
    fake_embedder: FakeEmbedder,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    owner = await register_user(client, unique_email, "知识库租户 A")
    headers_a = auth_headers(owner["tokens"]["access_token"])
    kb = await create_kb(client, headers_a)
    document = await upload_markdown(client, headers_a, kb["id"])
    _patch_worker_embedder(monkeypatch, fake_embedder)
    await process_document({}, document["id"], owner["user"]["tenant_id"])

    other = await register_user(client, f"kb-b-{uuid.uuid4().hex[:8]}@example.com", "知识库租户 B")
    headers_b = auth_headers(other["tokens"]["access_token"])

    bases_b = await client.get("/api/v1/knowledge/bases", headers=headers_b)
    assert bases_b.json() == []

    doc_b = await client.get(f"/api/v1/knowledge/documents/{document['id']}", headers=headers_b)
    assert doc_b.status_code == 404

    hits_b = await client.get("/api/v1/knowledge/search", params={"q": "退款"}, headers=headers_b)
    assert hits_b.json()["items"] == []


async def test_customer_cannot_manage_knowledge(
    client: AsyncClient, unique_email: str, fake_queue: FakeQueue
) -> None:
    owner = await register_user(client, unique_email)
    tenant_id = owner["user"]["tenant_id"]
    customer_email = f"customer-{uuid.uuid4().hex[:8]}@example.com"

    # 直接在 A 租户下建一个 customer 成员（注册接口只能开新租户）
    async with async_session_factory() as session:
        set_current_tenant(uuid.UUID(tenant_id))
        customer = User(
            id=uuid.uuid4(),
            email=customer_email,
            password_hash=hash_password(PASSWORD),
            display_name="只读客户",
            role="customer",
            tenant_id=uuid.UUID(tenant_id),
        )
        session.add(customer)
        await session.flush()
        session.add(
            TenantMember(tenant_id=uuid.UUID(tenant_id), user_id=customer.id, role="customer")
        )
        await session.commit()
    set_current_tenant(None)

    login = await client.post(
        "/api/v1/auth/login", json={"email": customer_email, "password": PASSWORD}
    )
    assert login.status_code == 200, login.text
    headers = auth_headers(login.json()["tokens"]["access_token"])

    created = await client.post(
        "/api/v1/knowledge/bases", json={"name": "越权建库"}, headers=headers
    )
    assert created.status_code == 403
    assert created.json()["code"] == "FORBIDDEN"

    # 读操作仍然允许（坐席要能查资料）
    listed = await client.get("/api/v1/knowledge/bases", headers=headers)
    assert listed.status_code == 200


async def test_document_reprocess_and_delete(
    client: AsyncClient,
    unique_email: str,
    fake_queue: FakeQueue,
    fake_embedder: FakeEmbedder,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = await register_user(client, unique_email)
    tenant_id = data["user"]["tenant_id"]
    headers = auth_headers(data["tokens"]["access_token"])
    kb = await create_kb(client, headers)
    document = await upload_markdown(client, headers, kb["id"])
    _patch_worker_embedder(monkeypatch, fake_embedder)
    await process_document({}, document["id"], tenant_id)

    reprocessed = await client.post(
        f"/api/v1/knowledge/documents/{document['id']}/reprocess", headers=headers
    )
    assert reprocessed.status_code == 202
    assert reprocessed.json()["status"] == "uploaded"
    await process_document({}, document["id"], tenant_id)

    detail = await client.get(f"/api/v1/knowledge/documents/{document['id']}", headers=headers)
    assert detail.json()["status"] == "done"

    deleted = await client.delete(f"/api/v1/knowledge/documents/{document['id']}", headers=headers)
    assert deleted.status_code == 204

    hits = await client.get("/api/v1/knowledge/search", params={"q": "退款"}, headers=headers)
    assert hits.json()["items"] == [], "删除文档必须级联清理 chunks"


def test_markdown_chunking_keeps_heading_metadata() -> None:
    sections = [
        Section(text="退款政策说明。" * 40, heading="退款流程"),
        Section(text="电子发票 24 小时内发送。", heading="发票"),
    ]
    chunks = chunk_sections(
        sections, strategy=CHUNK_STRATEGY_BY_HEADING, chunk_size=100, chunk_overlap=20
    )

    assert chunks[0].metadata["heading"] == "退款流程"
    assert all(len(chunk.content) <= 100 for chunk in chunks)
    assert chunks[-1].metadata["heading"] == "发票"
    assert [chunk.index for chunk in chunks] == list(range(len(chunks)))


def test_fixed_window_chunking_overlaps() -> None:
    chunks = chunk_sections(
        [Section(text="x" * 250)],
        strategy=CHUNK_STRATEGY_FIXED_WINDOW,
        chunk_size=100,
        chunk_overlap=20,
    )
    assert len(chunks) == 3
    assert chunks[0].content == "x" * 100
    assert chunks[1].content == "x" * 100


def test_markdown_and_docx_extraction() -> None:
    markdown = "# 售后\n\n## 退款\n\n七天内可退。\n\n## 发票\n\n次日发送。".encode()
    sections = extract_sections(markdown, ".md")
    assert [section.heading for section in sections] == ["售后 > 退款", "售后 > 发票"]

    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_heading("退款流程", level=1)
    docx.add_paragraph("用户可在七天内申请退款。")
    table = docx.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "时效"
    table.rows[0].cells[1].text = "3 个工作日"
    buffer = io.BytesIO()
    docx.save(buffer)

    docx_sections = extract_sections(buffer.getvalue(), ".docx")
    assert docx_sections[0].heading == "退款流程"
    assert "七天内" in docx_sections[0].text
    assert "3 个工作日" in docx_sections[0].text, "表格内容必须进入可检索文本"


def test_tokenize_for_search_strips_punctuation() -> None:
    tokens = tokenize_for_search("退款流程是什么？请在 3 天内联系 support@example.com！")
    assert "退款" in tokens.split()
    assert "3" in tokens.split()
    for punctuation in ("？", "！", "，", "@"):
        assert punctuation not in tokens
