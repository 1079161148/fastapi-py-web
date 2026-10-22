"""文档解析与分块（设计文档 5.2 / 9.4）。

解析与分块是纯函数（不碰数据库/网络），便于用固定样本做回归；
入库与向量化在 workers/tasks.py 编排。

支持格式：PDF（pypdf 抽取文本层）、DOCX（python-docx，含表格）、MD/TXT。

分块策略（设计文档 9.4 "策略可配置"）：
- by_heading：按标题层级切段（Markdown 的 #，DOCX 的 Heading 样式），
  段内按段落贪心打包到 chunk_size；超长段落按窗口切分并保留 overlap
- fixed_window：全文固定窗口 + 重叠，不保留标题信息

**语义分块未实现**：需要评估集先验证收益，当前没有可对比的数据，
按 YAGNI 推迟到评估集建立之后再决定（不写"看似支持"的空壳）。

扫描版 PDF 没有文本层，会解析出空内容 → 明确失败，不做 OCR（M3 范围外）。
"""

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any

import jieba
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.db.models.knowledge import CHUNK_STRATEGY_FIXED_WINDOW

# jieba 首次调用会加载词典并往 stderr 打日志，压到 WARNING 避免淹没应用日志
jieba.setLogLevel(logging.WARNING)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}

_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
# 检索 token 白名单：英文/数字词、CJK 连续块；标点与 emoji 一律丢弃
_TOKEN_PATTERN = re.compile(r"[0-9A-Za-z_]+|[\u4e00-\u9fff]+")


class DocumentParseError(AppError):
    def __init__(self, message: str) -> None:
        super().__init__(code="DOCUMENT_PARSE_FAILED", message=message, status_code=422)


@dataclass(frozen=True, slots=True)
class Section:
    """解析出的段落组：标题路径、页码（PDF）随 chunk 元数据保留以备引用溯源。"""

    text: str
    heading: str = ""
    page: int | None = None


@dataclass(frozen=True, slots=True)
class RawChunk:
    content: str
    index: int
    metadata: dict[str, Any] = field(default_factory=dict)


def extract_sections(data: bytes, suffix: str) -> list[Section]:
    """按格式抽取文本；不支持的后缀直接报错，不做猜测性解析。"""
    normalized = suffix.lower()
    if normalized == ".pdf":
        return _extract_pdf(data)
    if normalized == ".docx":
        return _extract_docx(data)
    if normalized in {".md", ".markdown"}:
        return _extract_markdown(data.decode("utf-8", errors="replace"))
    if normalized == ".txt":
        text = data.decode("utf-8", errors="replace")
        return [Section(text=text)] if text.strip() else []
    raise DocumentParseError(f"不支持的文件类型：{suffix}（支持 PDF / DOCX / Markdown / TXT）")


def build_chunks(
    data: bytes,
    suffix: str,
    *,
    strategy: str | None = None,
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
) -> list[RawChunk]:
    """解析 + 分块的完整入口（worker 与测试共用）。"""
    settings = get_settings()
    sections = extract_sections(data, suffix)
    if not sections:
        raise DocumentParseError("文档未解析出文本内容：可能是扫描版 PDF 或空文件")
    return chunk_sections(
        sections,
        strategy=strategy or settings.chunk_strategy,
        chunk_size=chunk_size or settings.chunk_size,
        chunk_overlap=chunk_overlap if chunk_overlap is not None else settings.chunk_overlap,
    )


def chunk_sections(
    sections: list[Section], *, strategy: str, chunk_size: int, chunk_overlap: int
) -> list[RawChunk]:
    if chunk_size <= 0:
        raise ValueError("chunk_size 必须为正数")
    overlap = max(0, min(chunk_overlap, chunk_size // 2))

    if strategy == CHUNK_STRATEGY_FIXED_WINDOW:
        joined = "\n\n".join(section.text for section in sections if section.text.strip())
        pieces = _window_split(joined, chunk_size, overlap)
        return [
            RawChunk(content=piece, index=index, metadata={"strategy": strategy})
            for index, piece in enumerate(pieces)
        ]

    chunks: list[RawChunk] = []
    for section in sections:
        if not section.text.strip():
            continue
        metadata: dict[str, Any] = {"strategy": strategy}
        if section.heading:
            metadata["heading"] = section.heading
        if section.page is not None:
            metadata["page"] = section.page
        for piece in _pack_paragraphs(section.text, chunk_size, overlap):
            chunks.append(RawChunk(content=piece, index=len(chunks), metadata=metadata))
    return chunks


def tokenize_for_search(text: str) -> str:
    """jieba 分词 → 空格串，作为 chunks.search_text（tsv 生成列的输入）。

    检索时查询串必须走同一函数，两侧分词口径不一致会导致关键词召回失效。
    """
    tokens: list[str] = []
    for raw in jieba.cut_for_search(text):
        token = raw.strip().lower()
        if token and _TOKEN_PATTERN.fullmatch(token):
            tokens.append(token)
    return " ".join(tokens)


def _extract_pdf(data: bytes) -> list[Section]:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf 对损坏文件抛多种异常，统一为可读错误
        raise DocumentParseError("PDF 解析失败：文件可能已损坏或被加密") from exc

    sections: list[Section] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            raise DocumentParseError(f"PDF 第 {page_number} 页解析失败") from exc
        if text:
            sections.append(Section(text=text, page=page_number))
    return sections


def _extract_docx(data: bytes) -> list[Section]:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("DOCX 解析失败：文件可能已损坏") from exc

    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            sections.append(Section(text="\n".join(buffer), heading=heading))
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.startswith("Heading") or style_name.startswith("标题"):
            flush()
            heading = text
        else:
            buffer.append(text)

    # 表格内容对客服问答价值高（价格/时效/规格），按行拼成可检索文本
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                buffer.append(" | ".join(cells))
    flush()
    return sections


def _extract_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    heading_stack: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        content = "\n".join(buffer).strip()
        if content:
            sections.append(Section(text=content, heading=" > ".join(heading_stack)))
        buffer.clear()

    for line in text.splitlines():
        match = _HEADING_PATTERN.match(line)
        if match:
            flush()
            level = len(match.group(1))
            heading_stack = heading_stack[: level - 1] + [match.group(2).strip()]
        else:
            buffer.append(line)
    flush()
    return sections


def _window_split(text: str, chunk_size: int, overlap: int) -> list[str]:
    content = text.strip()
    if not content:
        return []
    if len(content) <= chunk_size:
        return [content]

    pieces: list[str] = []
    step = max(chunk_size - overlap, 1)
    start = 0
    while start < len(content):
        piece = content[start : start + chunk_size].strip()
        if piece:
            pieces.append(piece)
        if start + chunk_size >= len(content):
            break
        start += step
    return pieces


def _pack_paragraphs(text: str, chunk_size: int, overlap: int) -> list[str]:
    """段落贪心打包：尽量保持语义完整，超长段落再退化为窗口切分。"""
    paragraphs = [line.strip() for line in re.split(r"\n+", text) if line.strip()]
    pieces: list[str] = []
    buffer: list[str] = []
    buffer_size = 0

    for paragraph in paragraphs:
        if len(paragraph) > chunk_size:
            if buffer:
                pieces.append("\n".join(buffer))
                buffer, buffer_size = [], 0
            pieces.extend(_window_split(paragraph, chunk_size, overlap))
            continue
        # +1 是换行符开销，避免打包结果刚好超出窗口
        if buffer and buffer_size + len(paragraph) + 1 > chunk_size:
            pieces.append("\n".join(buffer))
            buffer, buffer_size = [], 0
        buffer.append(paragraph)
        buffer_size += len(paragraph) + 1

    if buffer:
        pieces.append("\n".join(buffer))
    return pieces
